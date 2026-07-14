from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DOCS.mkdir(parents=True, exist_ok=True)

NAVY = "0A1633"
BLUE = "3F68FF"
DARK_BLUE = "1F4D78"
INK = "17223B"
MUTED = "66728A"
LINE = "DDE3EE"
LIGHT = "F3F6FB"
BLUE_LIGHT = "EEF2FF"
GREEN_LIGHT = "E2F8F0"
AMBER_LIGHT = "FFF2D7"
RED_LIGHT = "FFE6E8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size=7):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn("w:" + edge))
        if tag is None:
            tag = OxmlElement("w:" + edge)
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_fixed_table_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_run(run, size=10.5, bold=False, color=INK, name="Calibri", italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run(run, 8.5, color=MUTED)


def configure_document(doc, running_title):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run(hp.add_run(running_title.upper()), 8, True, MUTED, "Arial")
    rule = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), LINE)
    rule.append(bottom)
    hp._p.get_or_add_pPr().append(rule)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_run(fp.add_run("Moeen Ahmad Butt  |  SchoolTrust AI  |  Page "), 8.5, color=MUTED)
    add_page_field(fp)


def add_cover(doc, document_type, subtitle, metadata):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("CAPSTONE PROJECT  /  ENTERPRISE AI RAG SUPPORT")
    set_run(r, 9, True, BLUE, "Arial")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SCHOOLTRUST AI")
    set_run(r, 30, True, NAVY, "Arial")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Verified Parent Support and Knowledge Improvement System")
    set_run(r, 15, True, DARK_BLUE, "Arial")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run(subtitle)
    set_run(r, 11, color=MUTED, name="Arial")

    table = doc.add_table(rows=len(metadata), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(4.1)
    table.columns[1].width = Cm(12.0)
    set_fixed_table_layout(table)
    set_table_borders(table, color=LINE, size=5)
    for row, (label, value) in zip(table.rows, metadata):
        row.cells[0].width = Cm(4.1)
        row.cells[1].width = Cm(12.0)
        set_cell_shading(row.cells[0], BLUE_LIGHT)
        for cell in row.cells:
            set_cell_margins(cell, 100, 130, 100, 130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        set_run(row.cells[0].paragraphs[0].add_run(label), 9.5, True, DARK_BLUE)
        set_run(row.cells[1].paragraphs[0].add_run(value), 9.5, value.startswith("[ENTER"), INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(7)
    set_run(p.add_run(document_type.upper()), 11, True, BLUE, "Arial")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("One focused project. Approved sources. Human-controlled improvement."), 10.5, True, NAVY)
    doc.add_page_break()


def add_table(doc, headers, rows, widths=None, font_size=9.1, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_fixed_table_layout(table)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, 95, 110, 95, 110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        set_run(cell.paragraphs[0].add_run(text), 9, True, WHITE)
    for row_data in rows:
        body_row = table.add_row()
        prevent_row_split(body_row)
        cells = body_row.cells
        for idx, value in enumerate(row_data):
            cell = cells[idx]
            set_cell_margins(cell, 85, 110, 85, 110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run(p.add_run(str(value)), font_size, idx == 0, INK)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_note(doc, title, text, fill=BLUE_LIGHT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_fixed_table_layout(table)
    set_table_borders(table, color=accent, size=7)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 130, 150, 130, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(title), 10, True, DARK_BLUE)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(text), 9.7, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.space_after = Pt(4)
        set_run(p.add_run(item), 10.3)


def add_use_case(doc, uc):
    doc.add_heading(uc["id"] + " - " + uc["name"], level=2)
    rows = [
        ("Primary actor", uc["actor"]),
        ("Purpose", uc["purpose"]),
        ("Preconditions", uc["preconditions"]),
        ("Trigger", uc["trigger"]),
        ("Main flow", "\n".join(str(i + 1) + ". " + step for i, step in enumerate(uc["flow"]))),
        ("Alternative flows", "\n".join("- " + step for step in uc["alternatives"])),
        ("Postconditions", uc["postconditions"]),
        ("Related requirements", ", ".join(uc["requirements"])),
    ]
    add_table(doc, ["Field", "Specification"], rows, [3.6, 12.4], font_size=9.2, header_fill=DARK_BLUE)


def build_srs():
    doc = Document()
    configure_document(doc, "Software Requirements Specification  |  Version 1.0")
    add_cover(doc, "Software Requirements Specification", "Requirements-only specification for the assessed student implementation.", [
        ("Project title", "SchoolTrust AI"),
        ("Student name", "Moeen Ahmad Butt"),
        ("Registration ID", "[ENTER REGISTRATION ID]"),
        ("Case study", "Umeed Education System, Lahore"),
        ("Version / date", "Version 1.0  |  14 July 2026"),
        ("Live demo link", "[ENTER AFTER DEPLOYMENT]"),
        ("GitHub repository", "[ENTER IF AVAILABLE]"),
    ])

    doc.add_heading("Document Control", level=1)
    add_table(doc, ["Item", "Requirement"], [
        ("Document purpose", "Define the functional requirements, user roles, use cases, business rules and external interfaces for SchoolTrust AI."),
        ("Selected capstone", "Enterprise AI RAG Customer Support Platform."),
        ("Implementation boundary", "One school, one website assistant, one administrator dashboard and one n8n automation project."),
        ("Requirement language", "The word shall indicates a mandatory requirement. Should indicates a desirable requirement."),
        ("Prototype status", "Student capstone prototype requiring owner credentials and approved school content before live use."),
    ], [4.1, 11.9])

    doc.add_heading("System Boundary", level=1)
    add_table(doc, ["In scope", "Out of scope"], [
        ("Parent questions about approved public school information", "Private student academic, attendance or fee-ledger records"),
        ("English and Urdu question handling", "Voice calls and social-media omnichannel deployment"),
        ("Current-source retrieval and answer guardrails", "Independent AI authority to create or approve school policy"),
        ("Knowledge-gap recording and administrator approval", "Automatic publication without human review"),
        ("Regression evaluation of critical questions", "Production certification, formal penetration testing or legal compliance certification"),
    ], [8.0, 8.0])
    add_note(doc, "Core requirement", "A missing answer is not treated as a chatbot failure. It becomes a controlled knowledge-gap record for administrator review. The system shall prefer a safe escalation over an unsupported school fact.")

    doc.add_heading("1. Functional Requirements", level=1)
    fr_rows = [
        ("FR-01", "The system shall accept a parent question through the website webhook using JSON.", "Must", "API test"),
        ("FR-02", "The system shall reject questions shorter than 3 characters or longer than 1,000 characters.", "Must", "Boundary test"),
        ("FR-03", "The system shall detect whether the question is primarily English or Urdu.", "Must", "Language test"),
        ("FR-04", "The system shall classify the question into a supported topic such as fees, admissions, timings, contact, transport or general.", "Must", "Topic test"),
        ("FR-05", "The system shall mark fees, dates, timings, admissions, policies, payments, results and transport as critical topics.", "Must", "Rule test"),
        ("FR-06", "The system shall create a semantic embedding for every valid question.", "Must", "Execution log"),
        ("FR-07", "The system shall retrieve only approved knowledge with a non-expired effective window.", "Must", "Database test"),
        ("FR-08", "The system shall provide the language model with the question and retrieved approved sources only.", "Must", "Prompt inspection"),
        ("FR-09", "The generated response shall contain an answer mode, answer text, confidence score, source identifiers and an internal reason.", "Must", "Schema test"),
        ("FR-10", "The system shall verify that cited source identifiers exist in the retrieved source set.", "Must", "Guardrail test"),
        ("FR-11", "The system shall require a minimum confidence of 0.82 for a critical-topic verified answer.", "Must", "Threshold test"),
        ("FR-12", "The system shall require a minimum confidence of 0.72 for a non-critical verified answer.", "Must", "Threshold test"),
        ("FR-13", "The system shall request clarification when multiple valid student categories or cases require different answers.", "Must", "Scenario test"),
        ("FR-14", "The system shall safely escalate when approved information is absent, incomplete, conflicting or below the required threshold.", "Must", "Scenario test"),
        ("FR-15", "The system shall record repeated unsupported questions as grouped knowledge gaps with occurrence counts.", "Must", "Database test"),
        ("FR-16", "The system shall assign a higher initial priority to gaps involving critical topics.", "Must", "Rule test"),
        ("FR-17", "The system shall log verified, clarification and escalated conversations with topic, confidence and source identifiers.", "Must", "Log inspection"),
        ("FR-18", "The administrator dashboard shall display approved knowledge, open gaps, verified answers, escalations and seven-day test pass rate.", "Must", "UI/API test"),
        ("FR-19", "The system shall require an administrator token for dashboard and approval endpoints in the prototype.", "Must", "Authorization test"),
        ("FR-20", "An administrator shall be able to publish approved knowledge with title, category, content, version and effective dates.", "Must", "Approval test"),
        ("FR-21", "Publishing approved knowledge shall create its embedding and store an approval audit record.", "Must", "Database test"),
        ("FR-22", "Publishing an answer for a recorded gap shall resolve the related gap.", "Must", "Workflow test"),
        ("FR-23", "The system shall execute active regression questions manually and once per day.", "Must", "Schedule test"),
        ("FR-24", "The evaluation module shall store the generated answer, pass status, score and missing expected facts.", "Must", "Evaluation test"),
        ("FR-25", "The web interface shall provide suggested questions and visibly distinguish verified, clarification and escalated responses.", "Should", "UI inspection"),
        ("FR-26", "The web interface shall be responsive on desktop, tablet and mobile widths.", "Should", "Responsive test"),
    ]
    add_table(doc, ["ID", "Requirement", "Priority", "Verification"], fr_rows, [1.6, 10.2, 1.8, 2.4], font_size=8.7)

    doc.add_heading("2. User Roles", level=1)
    add_table(doc, ["Role", "Permissions", "Restrictions"], [
        ("Parent / visitor", "Ask public school-information questions; receive verified answers, clarification requests or escalation notices.", "Cannot view the administrator dashboard, approve knowledge or access private student records."),
        ("School administrator", "View knowledge health; review gaps; approve knowledge; set effective dates; resolve gaps; view test results.", "Cannot expose service-role or AI keys to the browser. Must approve only authorised school information."),
        ("Automation scheduler", "Run daily regression evaluation and store results.", "Cannot approve knowledge or alter policy. Operates only on active test cases and approved sources."),
    ], [3.2, 7.5, 5.3])

    doc.add_heading("3. Use Cases", level=1)
    use_cases = [
        {
            "id": "UC-01", "name": "Receive a verified answer", "actor": "Parent / visitor",
            "purpose": "Obtain an answer supported by current approved school knowledge.",
            "preconditions": "The workflow is active and relevant approved knowledge has an embedding.",
            "trigger": "The parent submits a valid question.",
            "flow": ["Validate the question.", "Detect language, topic and criticality.", "Create the embedding.", "Retrieve approved active sources.", "Generate a controlled response.", "Apply source and confidence guardrails.", "Log and return the verified answer."],
            "alternatives": ["If the source set is empty, continue to UC-03.", "If the answer does not cite a retrieved source, continue to UC-03."],
            "postconditions": "A verified conversation log exists and the parent receives the answer.",
            "requirements": ["FR-01", "FR-03", "FR-04", "FR-06", "FR-07", "FR-09", "FR-10", "FR-17"]
        },
        {
            "id": "UC-02", "name": "Request parent clarification", "actor": "Parent / visitor",
            "purpose": "Resolve ambiguity without selecting an unsupported category.",
            "preconditions": "Approved sources contain more than one applicable case.",
            "trigger": "The question lacks a required category, such as new admission versus promoted student.",
            "flow": ["Retrieve both applicable cases.", "Generate a precise clarification question or state both categories.", "Log the clarification response.", "Return it to the parent."],
            "alternatives": ["If the sources conflict rather than describe valid categories, continue to UC-03."],
            "postconditions": "The system waits for a more specific parent question.",
            "requirements": ["FR-13", "FR-17"]
        },
        {
            "id": "UC-03", "name": "Escalate and record a knowledge gap", "actor": "Parent / visitor",
            "purpose": "Avoid inventing school information and create a reviewable improvement task.",
            "preconditions": "The question is valid but cannot be answered safely.",
            "trigger": "No approved source, low confidence, conflict or incomplete information is detected.",
            "flow": ["Return a safe parent-facing notice.", "Normalise the topic.", "Create or increment the matching open gap.", "Assign priority.", "Log the escalated conversation."],
            "alternatives": ["If a similar open gap already exists, increment its occurrence count instead of creating a duplicate."],
            "postconditions": "The parent is not misled and an administrator can review the gap.",
            "requirements": ["FR-14", "FR-15", "FR-16", "FR-17"]
        },
        {
            "id": "UC-04", "name": "View knowledge health", "actor": "School administrator",
            "purpose": "Review the operational state of the assistant.",
            "preconditions": "A valid prototype administrator token is supplied.",
            "trigger": "The administrator loads dashboard data.",
            "flow": ["Validate the token.", "Calculate knowledge and conversation metrics.", "Load open gaps and recent test results.", "Return the dashboard response.", "Render live metrics."],
            "alternatives": ["Reject an absent or invalid token."],
            "postconditions": "No data is changed.",
            "requirements": ["FR-18", "FR-19"]
        },
        {
            "id": "UC-05", "name": "Approve and publish knowledge", "actor": "School administrator",
            "purpose": "Convert reviewed school information into an approved searchable source.",
            "preconditions": "The administrator has authoritative content and a valid token.",
            "trigger": "The approval form is submitted.",
            "flow": ["Validate mandatory fields.", "Create the content embedding.", "Store the source as approved.", "Record effective dates and version.", "Resolve the linked gap when present.", "Write an administrator audit record."],
            "alternatives": ["Reject content shorter than 20 characters.", "Reject an invalid administrator token."],
            "postconditions": "The approved source is available for retrieval and the change is auditable.",
            "requirements": ["FR-19", "FR-20", "FR-21", "FR-22"]
        },
        {
            "id": "UC-06", "name": "Run regression evaluation", "actor": "Automation scheduler or administrator",
            "purpose": "Verify critical answers after knowledge changes.",
            "preconditions": "Active regression tests and approved embedded knowledge exist.",
            "trigger": "The daily schedule fires or the manual trigger is executed.",
            "flow": ["Load active tests.", "Embed each question.", "Retrieve approved knowledge.", "Generate an answer using only that knowledge.", "Compare expected facts.", "Store score, status and missing facts."],
            "alternatives": ["A missing approved answer produces a failed test rather than an invented answer."],
            "postconditions": "Evaluation results are visible on the dashboard.",
            "requirements": ["FR-23", "FR-24"]
        },
    ]
    for uc in use_cases:
        add_use_case(doc, uc)

    doc.add_page_break()
    doc.add_heading("4. Business Rules", level=1)
    br_rows = [
        ("BR-01", "Only knowledge with status approved may support a parent answer."),
        ("BR-02", "A source with a future effective date or expired date shall not be retrieved."),
        ("BR-03", "Critical topics shall use the 0.82 verified-answer confidence threshold."),
        ("BR-04", "Non-critical topics shall use the 0.72 verified-answer confidence threshold."),
        ("BR-05", "A verified answer shall cite at least one identifier from the retrieved source set."),
        ("BR-06", "The assistant shall not invent fees, dates, timings, policies, addresses or availability."),
        ("BR-07", "Different valid categories shall be stated separately or clarified. They shall not be collapsed into one amount."),
        ("BR-08", "Conflicting or incomplete approved information shall result in escalation."),
        ("BR-09", "An administrator shall approve every knowledge publication."),
        ("BR-10", "The public web interface shall never receive the Supabase service-role key or AI API key."),
        ("BR-11", "Private student records are outside the public assistant's scope."),
        ("BR-12", "Repeated identical open gaps within the same topic shall increment occurrence count."),
        ("BR-13", "Critical-topic gaps shall initially receive high priority."),
        ("BR-14", "Approved content shall contain at least 20 characters."),
        ("BR-15", "The question length shall be between 3 and 1,000 characters."),
        ("BR-16", "The administrator demo token is a prototype control and shall be replaced by authenticated role-based access before production use."),
        ("BR-17", "The system is a controlled knowledge workflow, not a self-training model."),
    ]
    add_table(doc, ["ID", "Business rule"], br_rows, [1.8, 14.2], font_size=9.1)

    doc.add_heading("5. External Interfaces", level=1)
    doc.add_heading("5.1 User Interfaces", level=2)
    add_table(doc, ["Interface", "Requirement"], [
        ("Parent assistant", "Responsive chat view with question input, suggested questions, conversation stream and visible response status."),
        ("Administrator dashboard", "Metrics for approved knowledge, open gaps, verified answers and test pass rate; gap review; test results."),
        ("Approval dialog", "Fields for gap identifier, document title, category, approved content, effective dates and administrator name."),
    ], [4.0, 12.0])

    doc.add_heading("5.2 HTTP Interfaces", level=2)
    add_table(doc, ["Method / path", "Request", "Response / control"], [
        ("POST /webhook/schooltrust-ask", "question, session_id, optional user_contact", "Verified, clarification or escalated answer. Public prototype endpoint."),
        ("GET /webhook/schooltrust-admin-data", "Header: x-admin-token", "Dashboard metrics, gaps and evaluations. Administrator token required."),
        ("POST /webhook/schooltrust-admin-approve", "Header: x-admin-token plus approved source fields", "Published knowledge identifier and approval confirmation."),
    ], [5.2, 5.2, 5.6], font_size=8.8)

    doc.add_heading("5.3 Data and AI Interfaces", level=2)
    add_table(doc, ["External system", "Purpose", "Required configuration"], [
        ("Supabase PostgreSQL with pgvector", "Approved knowledge, embeddings, conversation logs, gaps, regression tests, evaluation results and audit actions.", "SUPABASE_URL and server-side service-role key."),
        ("OpenAI API", "Question and content embeddings; controlled answer generation; regression answer generation.", "OPENAI_API_KEY, chat model and embedding model."),
        ("n8n", "Webhook handling, branching, guardrails, database calls, approval automation and scheduled tests.", "Imported workflow plus protected server-side configuration."),
        ("Static web host / school website", "Parent assistant and administrator prototype interface.", "Production n8n base URL and non-secret demo token only."),
    ], [4.2, 7.0, 4.8], font_size=8.8)

    doc.add_heading("5.4 Data Contracts", level=2)
    add_table(doc, ["Object", "Required fields"], [
        ("Parent question", "question: string; session_id: string; user_contact: optional string"),
        ("Controlled answer", "success: boolean; status: verified | clarification | escalated; answer: string; session_id: string"),
        ("Approved knowledge", "document_title; category; content; version_no; status; effective_from; expires_at; embedding"),
        ("Knowledge gap", "question; normalized_topic; reason; occurrence_count; priority; status; timestamps"),
        ("Evaluation result", "regression_test_id; generated_answer; passed; score; missing_facts; notes"),
    ], [4.2, 11.8])

    doc.add_heading("6. Requirements Traceability", level=1)
    add_table(doc, ["Capability", "Functional requirements", "Use case", "Implemented workflow nodes"], [
        ("Verified parent answer", "FR-01 to FR-12", "UC-01", "Parent Question Webhook through Return Verified Answer"),
        ("Clarification", "FR-13, FR-17", "UC-02", "Needs Clarification, Log Clarification, Return Clarification"),
        ("Safe escalation and gap", "FR-14 to FR-17", "UC-03", "Record Knowledge Gap, Log Escalated Conversation"),
        ("Knowledge dashboard", "FR-18, FR-19", "UC-04", "Administrator Dashboard Webhook, Load Dashboard Data"),
        ("Administrator approval", "FR-19 to FR-22", "UC-05", "Administrator Approval Webhook through Log Administrator Action"),
        ("Regression evaluation", "FR-23, FR-24", "UC-06", "Daily / Manual Quality Test through Store Evaluation Result"),
        ("Responsive interface", "FR-25, FR-26", "UC-01, UC-04, UC-05", "web/index.html, web/styles.css, web/app.js"),
    ], [3.6, 4.0, 2.4, 6.0], font_size=8.5)

    doc.add_heading("7. Acceptance Requirements", level=1)
    add_table(doc, ["ID", "Acceptance condition"], [
        ("AR-01", "The annual-fund question returns both new-admission and promoted-student categories or asks a precise clarification."),
        ("AR-02", "A transport-fee question with no approved source returns a safe escalation and creates or increments a gap."),
        ("AR-03", "An invalid administrator token cannot load dashboard data or publish knowledge."),
        ("AR-04", "An administrator-approved source becomes retrievable only after its embedding is stored."),
        ("AR-05", "The daily or manual evaluation branch stores a result for every active regression test."),
        ("AR-06", "The web interface displays a different status label for verified, clarification and escalated responses."),
        ("AR-07", "The submitted n8n JSON imports and contains the parent, approval, dashboard and testing branches."),
    ], [1.8, 14.2])

    add_note(doc, "Submission note", "This SRS describes the included prototype package. Registration ID, live demo link and GitHub link remain editable because they were not provided at generation time.", fill=AMBER_LIGHT, accent="D79A35")

    out = DOCS / "SchoolTrust_AI_SRS.docx"
    doc.save(out)
    return out


def build_setup_guide():
    doc = Document()
    configure_document(doc, "Setup and Deployment Guide  |  SchoolTrust AI")
    add_cover(doc, "Setup and Deployment Guide", "A controlled setup sequence for the assessed prototype.", [
        ("Project", "SchoolTrust AI"),
        ("Student", "Moeen Ahmad Butt"),
        ("Target platform", "n8n + Supabase + OpenAI + static web interface"),
        ("Version / date", "Version 1.0  |  14 July 2026"),
        ("Estimated setup", "45 to 90 minutes after credentials are available"),
    ])
    doc.add_heading("1. Package Contents", level=1)
    add_table(doc, ["Path", "Purpose"], [
        ("n8n/SchoolTrust_AI_Verified_Parent_Support.json", "Importable automation with parent, administrator, dashboard and regression branches."),
        ("database/supabase_schema.sql", "Tables, vector function, dashboard function, gap grouping and test seed."),
        ("web/", "Responsive parent and administrator prototype."),
        ("data/", "Draft starter knowledge and regression-test reference CSV files."),
        ("docs/", "SRS, setup guide, demo script and checklist."),
    ], [7.0, 9.0])

    doc.add_heading("2. Prerequisites", level=1)
    add_bullets(doc, [
        "An n8n instance that permits webhook, schedule, code and HTTP Request nodes.",
        "A Supabase project with permission to enable pgvector.",
        "An OpenAI API key with embedding and chat-model access.",
        "A static hosting location or access to the existing school website.",
        "Administrator-approved school information. Do not publish unreviewed draft data.",
    ])

    doc.add_heading("3. Configure Supabase", level=1)
    add_bullets(doc, [
        "Open the Supabase SQL editor.",
        "Review database/supabase_schema.sql. Use a new project when possible.",
        "Execute the script once.",
        "Confirm that six tables exist: knowledge_chunks, conversation_logs, knowledge_gaps, regression_tests, evaluation_results and admin_actions.",
        "Confirm that match_school_knowledge, record_knowledge_gap and schooltrust_admin_dashboard functions exist.",
        "Keep the service-role key server-side. Never place it in the web folder.",
    ])
    add_note(doc, "Important", "The included starter knowledge CSV is marked draft. Publish each reviewed entry through the administrator approval endpoint so its embedding and approval metadata are created.", fill=AMBER_LIGHT, accent="D79A35")

    doc.add_heading("4. Import and Configure n8n", level=1)
    add_bullets(doc, [
        "Import n8n/SchoolTrust_AI_Verified_Parent_Support.json.",
        "Keep the workflow inactive during configuration.",
        "Set SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY, OPENAI_API_KEY, OPENAI_CHAT_MODEL, OPENAI_EMBEDDING_MODEL and SCHOOLTRUST_ADMIN_TOKEN.",
        "If custom environment values are unavailable, replace the expressions with encrypted n8n credentials. Do not hardcode secrets in the browser.",
        "Execute each HTTP node once during testing and check for a 2xx response.",
        "Run the Manual Quality Test trigger after approved knowledge exists.",
        "Activate the workflow only after the ask, dashboard and approval production webhooks respond correctly.",
    ])
    add_table(doc, ["Variable", "Example / purpose"], [
        ("SUPABASE_URL", "https://YOUR_PROJECT.supabase.co"),
        ("SUPABASE_SERVICE_ROLE_KEY", "Server-side database access. Treat as secret."),
        ("OPENAI_API_KEY", "Server-side AI access. Treat as secret."),
        ("OPENAI_CHAT_MODEL", "gpt-4.1-mini or another JSON-capable approved model."),
        ("OPENAI_EMBEDDING_MODEL", "text-embedding-3-small, matching vector(1536)."),
        ("SCHOOLTRUST_ADMIN_TOKEN", "A long random token used only for the prototype administrator endpoints."),
    ], [5.0, 11.0])

    doc.add_heading("5. Publish Initial Knowledge", level=1)
    p = doc.add_paragraph()
    set_run(p.add_run("Use the Administrator Approval Webhook. Required body fields:"), 10.3, True)
    add_table(doc, ["Field", "Example"], [
        ("document_title", "Annual Fund Policy"),
        ("category", "fees"),
        ("content", "The exact administrator-approved answer."),
        ("actor", "School Administrator"),
        ("effective_from", "2026-01-01"),
        ("expires_at", "Optional"),
        ("gap_id", "Optional UUID when resolving a gap"),
    ], [5.0, 11.0])
    add_note(doc, "Recommended first records", "School profile, official contact, classes offered, annual-fund categories and one active admissions policy. Keep time-sensitive notices in separate entries with expiry dates.")

    doc.add_heading("6. Configure the Web Interface", level=1)
    add_bullets(doc, [
        "Open web/config.js.",
        "Replace https://YOUR-N8N-DOMAIN with the n8n production base URL.",
        "Replace the demo administrator token with the same SCHOOLTRUST_ADMIN_TOKEN value.",
        "Deploy the web folder to a static host or the existing website.",
        "For production, replace the browser-held administrator token with a real authenticated administrator session.",
    ])

    doc.add_heading("7. Required Test Sequence", level=1)
    add_table(doc, ["Test", "Expected result"], [
        ("Ask the annual-fund question", "Both categories are stated or a precise clarification is requested."),
        ("Ask an unsupported transport-fee question", "Safe escalation plus an open knowledge gap."),
        ("Load dashboard without token", "Request rejected."),
        ("Load dashboard with valid token", "Live metrics, gaps and test results returned."),
        ("Approve the transport answer", "Approved knowledge stored, embedding created and related gap resolved."),
        ("Run manual quality test", "One evaluation result stored for each active test."),
        ("Ask the transport question again", "Verified answer returned if the new source meets retrieval and confidence rules."),
    ], [6.0, 10.0])

    doc.add_heading("8. Troubleshooting", level=1)
    add_table(doc, ["Problem", "Check"], [
        ("Web interface says setup required", "Confirm web/config.js contains the production n8n base URL."),
        ("401 or unauthorized administrator request", "Confirm x-admin-token matches SCHOOLTRUST_ADMIN_TOKEN."),
        ("Supabase RPC not found", "Run the complete SQL schema and refresh the Supabase API cache."),
        ("No knowledge retrieved", "Confirm status is approved, dates are active and embedding is not null."),
        ("Vector dimension error", "Use text-embedding-3-small with vector(1536), or update both schema and model together."),
        ("Model response cannot be parsed", "Use a chat model supporting JSON-object response format."),
        ("Quality tests all fail", "Publish approved source records before running the test branch."),
    ], [6.0, 10.0])

    out = DOCS / "Setup_and_Deployment_Guide.docx"
    doc.save(out)
    return out


def build_demo_script():
    doc = Document()
    configure_document(doc, "Demo Presentation Script  |  SchoolTrust AI")
    add_cover(doc, "Demo Presentation Script", "A six-minute capstone walkthrough with one complete improvement loop.", [
        ("Project", "SchoolTrust AI"),
        ("Student", "Moeen Ahmad Butt"),
        ("Recommended duration", "6 to 8 minutes"),
        ("Required browser tabs", "Web interface, n8n workflow and Supabase tables"),
        ("Version / date", "Version 1.0  |  14 July 2026"),
    ])
    doc.add_heading("1. Opening - 40 seconds", level=1)
    add_note(doc, "Say this", "My project is SchoolTrust AI, a verified parent-support and knowledge-improvement system. The problem is not simply answering school questions. The real problem is preventing outdated or incomplete answers and improving the knowledge base when the system cannot answer safely.")

    doc.add_heading("2. Show the Real Problem - 45 seconds", level=1)
    add_bullets(doc, [
        "Explain that a normal chatbot may give only one annual-fund amount.",
        "State that the correct answer depends on whether the student is a new admission or a promoted student.",
        "Explain that generic fallback responses also waste the knowledge already available.",
    ])

    doc.add_heading("3. Verified Answer - 60 seconds", level=1)
    add_table(doc, ["Action", "What to explain"], [
        ("Ask: What is the annual fund?", "The workflow classifies fees as critical and applies the higher confidence requirement."),
        ("Show the response", "The assistant states both categories or asks which category applies."),
        ("Point to the status label", "Verified means the answer cited an active approved source and passed guardrails."),
    ], [6.2, 9.8])

    doc.add_heading("4. Knowledge Failure Recovery - 90 seconds", level=1)
    add_table(doc, ["Action", "What to explain"], [
        ("Ask: What is the current transport fee?", "No approved transport source exists."),
        ("Show safe escalation", "The AI does not invent an amount."),
        ("Open Gap Review", "The unsupported question is recorded with topic, reason, priority and occurrence count."),
        ("Ask again if useful", "A repeated identical question increments the open gap instead of creating uncontrolled duplicates."),
    ], [6.2, 9.8])

    doc.add_heading("5. Human Approval - 75 seconds", level=1)
    add_table(doc, ["Action", "What to explain"], [
        ("Open Review and Approve", "The administrator enters the exact authorised answer and effective dates."),
        ("Publish", "n8n creates the embedding, stores the approved record, resolves the gap and writes an audit action."),
        ("Run Manual Quality Test", "Critical questions are retested after the knowledge change."),
    ], [6.2, 9.8])

    doc.add_heading("6. Close the Loop - 45 seconds", level=1)
    add_bullets(doc, [
        "Ask the same transport-fee question again.",
        "Show that the system now retrieves the administrator-approved source.",
        "Open Knowledge Health and show test pass rate and gap count.",
        "State that this is the distinctive part: question, verification, safe failure, human approval and retesting.",
    ])

    doc.add_heading("7. Technical Summary - 45 seconds", level=1)
    add_table(doc, ["Layer", "Implementation"], [
        ("Interface", "Responsive HTML, CSS and JavaScript parent and administrator views."),
        ("Automation", "One n8n workflow with 38 nodes and four logical modules."),
        ("Data", "Supabase PostgreSQL, pgvector, RLS and server-side RPC functions."),
        ("AI", "Embeddings for retrieval and a controlled JSON response for answer generation."),
        ("Safety", "Approved-source filtering, effective dates, source validation, critical thresholds and human approval."),
    ], [4.0, 12.0])

    doc.add_heading("8. Likely Assessor Questions", level=1)
    add_table(doc, ["Question", "Answer"], [
        ("Is this just a chatbot?", "No. The chatbot is one interface. The core project is the answer-assurance and knowledge-improvement automation."),
        ("Does it train itself?", "No. It records gaps and proposes a workflow. An authorised administrator approves every new source."),
        ("Why use AI?", "AI handles language understanding, semantic retrieval and controlled response generation. Safety decisions remain rule-based."),
        ("What happens if sources conflict?", "The workflow clarifies valid categories or escalates incompatible information. It does not choose one silently."),
        ("Can it access student records?", "Not in this project. Public information and private student systems are intentionally separated."),
        ("What makes it student-feasible?", "It is limited to one school, three webhooks, one database and a controlled set of approved documents and tests."),
    ], [6.0, 10.0])
    add_note(doc, "Final sentence", "SchoolTrust AI does not only answer what the school already knows. It shows the school what its support system still needs to know, and improves only after human approval.", fill=GREEN_LIGHT, accent="12A879")

    out = DOCS / "Demo_Presentation_Script.docx"
    doc.save(out)
    return out


def build_market_brief():
    doc = Document()
    configure_document(doc, "Market Validation and Innovation Brief  |  SchoolTrust AI")
    add_cover(doc, "Market Validation and Innovation Brief", "Evidence for the problem, existing baseline and selected innovation gap.", [
        ("Project", "SchoolTrust AI"),
        ("Student", "Moeen Ahmad Butt"),
        ("Research method", "Official product and documentation review plus case-study observation"),
        ("Case-study evidence", "Umeed school chatbot fee and fallback-answer failure"),
        ("Version / date", "Version 1.0  |  14 July 2026"),
    ])

    doc.add_heading("1. Validated Problem", level=1)
    add_table(doc, ["Evidence", "Finding", "Project implication"], [
        ("Umeed case study", "The existing school assistant had a knowledge base but still returned a generic office-contact fallback for questions whose answers existed.", "Retrieval success and answer quality must be tested, not assumed."),
        ("Annual-fund case", "One generic amount was incomplete because new admissions and promoted students use different categories.", "Critical answers require category clarification and stronger guardrails."),
        ("Changing school information", "Fees, timings, notices and admissions information can become stale.", "Approved sources require effective and expiry dates."),
        ("Repeated parent questions", "Unsupported questions are usually answered manually but not converted into structured improvements.", "Unresolved questions should become prioritised knowledge gaps."),
    ], [3.2, 6.4, 6.4], font_size=8.8)

    doc.add_heading("2. Existing Market Baseline", level=1)
    add_table(doc, ["Source", "Existing capability", "What remains to demonstrate"], [
        ("Zendesk AI Knowledge Base", "Identifies content gaps, detects outdated resources and suggests new topics from support demand.", "A student implementation still needs a controlled end-to-end workflow that works on school policy data."),
        ("Ada knowledge guidance", "Treats unresolved AI conversations as evidence of missing knowledge and emphasises authorised sources and freshness.", "The capstone must turn this principle into visible gap, approval and effective-date controls."),
        ("Intercom content recommendations", "Uses unresolved conversations to recommend knowledge improvements ranked by impact.", "The school prototype needs a simpler, believable administrator approval loop."),
        ("n8n evaluation documentation", "Supports running AI workflow test datasets and evaluation logic.", "The capstone should prove critical school answers through regression tests."),
    ], [3.6, 6.2, 6.2], font_size=8.6)

    doc.add_heading("3. Gap Selected for the Capstone", level=1)
    add_note(doc, "Selected gap", "Small schools need an affordable support workflow that verifies critical public information, records unresolved questions, requires administrator approval, and retests important answers after every knowledge improvement.")
    add_table(doc, ["Ordinary RAG chatbot", "SchoolTrust AI improvement"], [
        ("Retrieves text and generates an answer.", "Filters by approval status and effective date before generation."),
        ("Uses one confidence approach for every question.", "Applies a stronger threshold to fees, dates, timings and policies."),
        ("Returns a generic fallback.", "Creates or increments a structured knowledge-gap record."),
        ("Updates knowledge manually without proof.", "Records administrator approval and runs regression questions."),
        ("May silently select one conflicting value.", "States valid categories separately, asks clarification or escalates."),
        ("Focuses only on conversation.", "Provides a knowledge-health dashboard and test results."),
    ], [7.4, 8.6])

    doc.add_heading("4. Innovation Claim", level=1)
    add_note(doc, "Honest positioning", "SchoolTrust AI does not claim that RAG, content-gap detection or regression testing are individually new. Its capstone innovation is the school-specific combination of critical-fact assurance, date-valid sources, human-approved knowledge-gap recovery and automatic retesting in one student-feasible workflow.", fill=GREEN_LIGHT, accent="12A879")
    add_table(doc, ["Innovation dimension", "Implementation evidence"], [
        ("Problem specificity", "Different school categories can produce different valid fee answers."),
        ("Safety design", "Source identifier validation plus 0.82 critical-answer threshold."),
        ("Human control", "Administrator token, approval form, audit log and no automatic policy authority."),
        ("Learning loop", "Repeated gap grouping, approved update, re-embedding and regression evaluation."),
        ("Commercial relevance", "Reduces repetitive office questions while exposing missing or stale support content."),
        ("Student feasibility", "One school, three webhooks, one database, one static interface and 38 n8n nodes."),
    ], [4.5, 11.5])

    doc.add_heading("5. Research Sources", level=1)
    add_table(doc, ["Organisation", "Source URL"], [
        ("Zendesk", "https://www.zendesk.com/service/help-center/ai-knowledge-base/"),
        ("Ada", "https://www.ada.cx/blog/crafting-a-smarter-knowledge-base-the-key-to-ai-customer-service-software-that-works/"),
        ("Intercom", "https://www.intercom.com/help/en/articles/11394959-use-ai-powered-content-recommendations-to-improve-fin"),
        ("n8n", "https://docs.n8n.io/build/integrate-ai/test-and-improve-ai-workflows/understand-why-to-test"),
    ], [3.2, 12.8], font_size=8.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    set_run(p.add_run("Research limitation: "), 9.2, True, DARK_BLUE)
    set_run(p.add_run("This review establishes a credible product gap for a capstone. It is not a patent search and does not support a claim of worldwide novelty."), 9.2, color=MUTED)

    out = DOCS / "Market_Validation_and_Innovation_Brief.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    paths = [build_srs(), build_setup_guide(), build_demo_script(), build_market_brief()]
    for path in paths:
        print(path)
